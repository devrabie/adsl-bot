import os
import datetime
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_dsl_report_excel(lines_data: List[Dict[str, Any]], filepath: str) -> str:
    """
    Generates a professionally styled Excel report (.xlsx) for ADSL lines.

    lines_data item schema:
    {
        "id": int,
        "phone_number": str,
        "total_gb": float,
        "used_gb": float,
        "remaining_gb": float,
        "net_consumption_24h": float,
        "expiry_date": str,
        "days_left": int,
        "status_str": str ("حرج", "تنبيه", "مستقر", "خطأ"),
        "last_error": Optional[str]
    }
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "تقرير خطوط الإنترنت"

    # Set Right-To-Left (RTL) layout
    ws.views.sheetView[0].rightToLeft = True

    # Styling definitions
    font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=11, bold=False, color="000000")
    font_summary = Font(name="Calibri", size=11, bold=True, color="1F4E78")

    fill_title = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fill_header = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    fill_summary = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    # Soft Highlight fills for status
    fill_critical = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Soft red/orange
    fill_warning = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # Soft yellow
    fill_stable = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")   # Soft green
    fill_error = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid")    # Soft dark red

    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border_side = Side(border_style="thin", color="D9D9D9")
    border_thin = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # Title Row
    ws.merge_cells("A1:K1")
    title_cell = ws["A1"]
    title_cell.value = f"تقرير مراقبة خطوط الإنترنت - يمن نت ({datetime.date.today().strftime('%Y-%m-%d')})"
    title_cell.font = font_title
    title_cell.fill = fill_title
    title_cell.alignment = align_center
    ws.row_dimensions[1].height = 40

    # Headers
    headers = [
        "معرف الخط",
        "رقم الخط",
        "اسم المشترك",
        "الباقة",
        "الرصيد الكلي (GB)",
        "المستهلك (GB)",
        "المتبقي (GB)",
        "استهلاك 24 ساعة (GB)",
        "تاريخ الانتهاء",
        "الأيام المتبقية",
        "الحالة العامة"
    ]

    ws.row_dimensions[3].height = 28
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.value = header_title
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_thin

    # Fill Data
    start_row = 4
    total_remaining_sum = 0.0
    total_24h_sum = 0.0

    for idx, data in enumerate(lines_data):
        current_row = start_row + idx
        ws.row_dimensions[current_row].height = 24

        rem_gb = data.get("remaining_gb", 0.0)
        days = data.get("days_left", 0)
        status_str = data.get("status_str", "مستقر")
        last_err = data.get("last_error")

        total_remaining_sum += rem_gb
        total_24h_sum += data.get("net_consumption_24h", 0.0)

        row_values = [
            data.get("id"),
            data.get("phone_number"),
            data.get("subscriber_name") or "-",
            data.get("package_name") or "-",
            data.get("total_gb", 0.0),
            data.get("used_gb", 0.0),
            rem_gb,
            data.get("net_consumption_24h", 0.0),
            data.get("expiry_date", "-"),
            days,
            status_str if not last_err else f"خطأ: {last_err}"
        ]

        # Determine soft row fill
        row_fill = fill_stable
        if last_err:
            row_fill = fill_error
        elif rem_gb < 10 or days < 5 or status_str == "حرج":
            row_fill = fill_critical
        elif rem_gb < 50 or days < 10 or status_str == "تنبيه":
            row_fill = fill_warning

        for col_num, val in enumerate(row_values, 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.value = val
            cell.font = font_data
            cell.fill = row_fill
            cell.alignment = align_center
            cell.border = border_thin

    # Summary Row
    summary_row = start_row + len(lines_data) + 1
    ws.row_dimensions[summary_row].height = 26

    ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=6)
    sum_label_cell = ws.cell(row=summary_row, column=1)
    sum_label_cell.value = f"الإجمالي الكلي ({len(lines_data)} خطوط):"
    sum_label_cell.font = font_summary
    sum_label_cell.fill = fill_summary
    sum_label_cell.alignment = align_right

    for c in range(1, 7):
        ws.cell(row=summary_row, column=c).fill = fill_summary
        ws.cell(row=summary_row, column=c).border = border_thin

    # Total Remaining GB cell
    cell_rem_sum = ws.cell(row=summary_row, column=7)
    cell_rem_sum.value = round(total_remaining_sum, 2)
    cell_rem_sum.font = font_summary
    cell_rem_sum.fill = fill_summary
    cell_rem_sum.alignment = align_center
    cell_rem_sum.border = border_thin

    # Total 24h Consumption cell
    cell_24h_sum = ws.cell(row=summary_row, column=8)
    cell_24h_sum.value = round(total_24h_sum, 2)
    cell_24h_sum.font = font_summary
    cell_24h_sum.fill = fill_summary
    cell_24h_sum.alignment = align_center
    cell_24h_sum.border = border_thin

    for c in range(9, 12):
        cell_empty = ws.cell(row=summary_row, column=c)
        cell_empty.fill = fill_summary
        cell_empty.border = border_thin

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            # approximate width for Arabic and UTF-8 characters
            len_val = len(val_str)
            if len_val > max_len:
                max_len = len_val
        ws.column_dimensions[col_letter].width = max(max_len + 5, 15)

    wb.save(filepath)
    return filepath
